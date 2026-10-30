"""THE TAKEOFF EXPORT (T5): a per-trade Excel of what the plans count, where
every figure is one the plan gate binds, and every figure it does not is
listed with its reason.

WHERE THE FIGURES COME FROM. Nothing here counts. The candidates are read
from the same books the chat's count answers rest on (plan_search.glyph_books
for the mechanical symbols, opening_books for door and window tags,
lib.plan_areas rows for areas), with the same helpers the evidence render
uses (_floor_count, _building_count, _qty_check). Then each candidate is put
to the gate as a plain sentence - "There are 3 PTAC-1." - with
plan_search.answer_is_grounded on the very records it came from. Only a
sentence the gate binds is exported as a figure. A line the books withhold,
or the gate refuses, goes to "Not verified / refused" with its reason; a
takeoff never shows a number the chat would refuse to say.

THE TRADES (Roy, 2026-10-28): doors and windows, equipment, roof drains,
areas; walls and fixtures join as their families land. A family is a roof
drain when its name says DRAIN or its tags are all RD-n; every other
mechanical or plan-only family is equipment.

EVERY LINE CARRIES ITS SOURCE: the sheets, the file and page of each, and
the drawing set it is per. Where the project has two drawing sets, every
line is per set (each set is counted on its own) - a door or window, and
since ruling 72 (2026-10-28) a mechanical or plan-only symbol too - and its
sentence names the set, as the gate requires of a figure where the sets may
differ. (The mechanical lines were read off the totals book, which counts
each sheet's symbols from one set, the filed where it has them: the owner's
set's 9 WH-1 appeared on no line.)

Pure: no database, no network. server.get_project_takeoff_xlsx loads the
evidence, calls build(), and writes() the workbook.
"""

from __future__ import annotations

import io
import re
from typing import Any, Dict, Iterable, List, NamedTuple, Optional, Sequence, Tuple

from lib import plan_areas as _areas
from lib import plan_search as S
from lib import plan_sources as _PSRC

DOORS_WINDOWS = "Doors & windows"
EQUIPMENT = "Equipment"
ROOF_DRAINS = "Roof drains"
AREAS = "Areas"
#: The sheet order of the workbook; a trade with no lines still gets its sheet
#: (an empty sheet says "none found", not "not looked for").
TRADES = (DOORS_WINDOWS, EQUIPMENT, ROOF_DRAINS, AREAS)
REFUSED_SHEET = "Not verified - refused"
#: A sheet's legend symbol count (lib.plan_legend_symbols, ruling 77): listed,
#: never a figure, until the sub-bid benchmark has measured the matcher.
LEGEND_SYMBOLS = "Legend symbols"
LEGEND_NOT_VERIFIED = ("legend symbol count: not verified until the sub-bid benchmark "
                       "(drawn symbols, sheet scope)")

BUILDING, FLOOR = "Building", "Floor"


class Line(NamedTuple):
    trade: str
    item: str                   # the schedule (family) or the area's name
    tag: str                    # "PTAC-1", "type F", "2A"
    scope: str                  # Building | Floor | Apartment | Room
    level: str                  # the floor as printed ("FIRST FLOOR PLAN"), or ""
    qty: Optional[float]        # None on a refused line
    unit: str                   # "ea" | "sq ft"
    set_label: str              # the drawing set it is per
    sheets: str                 # "A-100.00, M-100.00"
    source: str                 # "A-100.00: ARCH.pdf p4; M-100.00: MH.pdf p5"
    check: str                  # the schedule's QTY cross-check, or the area check
    sentence: str               # what the gate was asked to bind
    reason: str = ""            # why it is not verified ("" on a verified line)


class Takeoff(NamedTuple):
    lines: List[Line]           # verified
    refused: List[Line]

    def by_trade(self) -> Dict[str, List[Line]]:
        out: Dict[str, List[Line]] = {t: [] for t in TRADES}
        for ln in self.lines:
            out.setdefault(ln.trade, []).append(ln)
        return out


# ── naming ─────────────────────────────────────────────────────────────────

_ORDINAL_WORD = {n: w for w, n in S._ORDINALS.items()}


def _ordinal(n: int) -> str:
    if n in _ORDINAL_WORD:
        return _ORDINAL_WORD[n]
    suf = "th" if 10 <= n % 100 <= 20 else {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
    return f"{n}{suf}"


def _floor_subject(book: S.GlyphBook, k: str) -> str:
    """How a sentence names floor k so the gate reads it as k and only k: its
    level ("the first floor", "the roof") where that level is this floor's
    alone, else its sheet."""
    for name in sorted(book.levels.get(k) or (), key=str):
        if {kk for kk, ns in book.levels.items() if name in ns} != {k}:
            continue
        if isinstance(name, int):
            return f"the {_ordinal(name)} floor"
        if S._NAMED_LEVEL.fullmatch(str(name)):
            return f"the {name}"
    sheets = book.sheets.get(k) or []
    return f"sheet {sheets[0]}" if sheets else ""


def _trade(fam: S.GlyphFamily) -> str:
    tags = [str(t) for t in fam.tags or []]
    if re.search(r"\bDRAINS?\b", fam.name or "", re.I) or (
            tags and all(re.fullmatch(r"R\.?D-?\d*[A-Z]?", t, re.I) for t in tags)):
        return ROOF_DRAINS
    return EQUIPMENT


def _opening_noun(fam: S.GlyphFamily, n: int) -> str:
    noun = "window" if re.search(r"WINDOW", fam.name or "", re.I) else "door"
    return noun if n == 1 else noun + "s"


def _thing(fam: S.GlyphFamily, tag: str, n: int, opening: bool) -> str:
    if opening:
        return f"{S._type_name(tag)} {_opening_noun(fam, n)}"
    return tag


def _set_phrase(label: str, sets: int) -> str:
    return f" per {label}" if sets > 1 and label else ""


def _sentence(subject: str, n: int, thing: str, set_phrase: str) -> str:
    if not subject:
        return f"There {'is' if n == 1 else 'are'} {n} {thing}{set_phrase}."
    return f"{subject[0].upper()}{subject[1:]} has {n} {thing}{set_phrase}."


# ── sources ────────────────────────────────────────────────────────────────

def _places(r: Dict[str, Any]) -> Iterable[Tuple[str, str, Any, Optional[str]]]:
    """(sheet, file_id, page, file_name) of a glyph row: the sheet the symbol
    is on and, for a mechanical symbol, the architectural sheet it is placed
    on."""
    yield (str(r.get("sheet_number") or ""), str(r.get("file_id") or ""),
           r.get("page_number"), r.get("file_name"))
    a = (r.get("payload") or {}).get("arch") if isinstance(r.get("payload"), dict) else None
    if isinstance(a, dict) and (a.get("sheet_number") or a.get("page_id")) \
            and a.get("page_id") != r.get("page_id"):
        yield (str(a.get("sheet_number") or ""), str(a.get("file_id") or ""),
               a.get("page_number"), a.get("file_name"))


def _source(rows: Sequence[Dict[str, Any]], sources: Dict[str, dict]) -> Tuple[str, str]:
    """("A-100.00, M-100.00", "A-100.00: ARCH.pdf p4; M-100.00: MH.pdf p5")."""
    seen: Dict[Tuple[str, str], str] = {}
    for r in rows:
        for sheet, fid, page, fname in _places(r):
            name = ((sources.get(fid) or {}).get("file_name") or fname or "").strip()
            where = " ".join(x for x in (name, f"p{page}" if page is not None else "") if x)
            # A replacing page that prints no number is cited as the number it
            # was given, and says so (ruling 136): 'A-105.01 (gas-change set,
            # unnumbered) "Roof Plan"'.
            seen.setdefault((sheet, where), _PSRC.tag_sheet(sheet, sources.get(fid), page) if sheet else sheet)
    keys = sorted(seen, key=lambda k: (k[0] == "", k[0], k[1]))
    sheets = ", ".join(dict.fromkeys(seen[k] for k in keys if k[0]))
    src = "; ".join(f"{seen[k]}: {k[1]}" if k[0] and k[1] else (seen[k] or k[1]) for k in keys if k[0] or k[1])
    return sheets, src


def sheet_numbers(sheets: str) -> List[str]:
    """The bare sheet numbers of a `sheets` cell: "A-100.00, A-105.01
    (gas-change set, unnumbered) "Roof Plan"" -> ["A-100.00", "A-105.01"]. The
    tag holds a comma, so the cell is not split on ", " (the review queue's
    items name a sheet by its number)."""
    plain = _PSRC.without_tags(sheets)
    return [x.strip() for x in plain.split(",") if x.strip()]


def _row_set(r: Dict[str, Any], sources: Dict[str, dict]) -> str:
    """The drawing set a glyph row is per: its architectural sheet's."""
    p = r.get("payload") if isinstance(r.get("payload"), dict) else {}
    a = p.get("arch") if isinstance(p.get("arch"), dict) else {}
    src = sources.get(str(a.get("file_id") or ""))
    if src and src.get("set_name"):
        return src["set_name"]
    t = a.get("drawing_track") or p.get("track")
    if a.get("drawing_track_label"):
        return a["drawing_track_label"]
    return S._FILED_LABEL if t in (None, S._FILED) else f"the {t} set"


# ── why a figure is not verified ───────────────────────────────────────────

def _refusal_text(rows: Sequence[Dict[str, Any]]) -> str:
    parts = []
    for r in rows:
        p = r.get("payload") if isinstance(r.get("payload"), dict) else {}
        why = p.get("reason") or p.get("how") or ""
        s = str(r.get("sheet_number") or (p.get("arch") or {}).get("sheet_number") or "a sheet")
        parts.append(f"{s}" + (f" ({why})" if why else ""))
    return ", ".join(dict.fromkeys(parts))


def _floor_why(book: S.GlyphBook, fam: S.GlyphFamily, k: str) -> str:
    if fam.why:
        return fam.why
    if k in book.invariant:
        return book.invariant[k]
    ref = [r for r in S._blocking_refusals(book) if S._row_floor(r) == k]
    if ref:
        return "a sheet on this floor could not be read: " + _refusal_text(ref)
    if book.opening and any(S._unread(r) and S._row_floor(r) == k for r in fam.rows):
        return "a tag or symbol on this floor could not be read"
    return "nothing located"


def _building_why(book: S.GlyphBook, fam: S.GlyphFamily, tags: Sequence[str],
                  n: int) -> str:
    if fam.why:
        return fam.why
    ref = S._blocking_refusals(book)
    if ref:
        return "sheets that could carry these symbols could not be read: " + _refusal_text(ref)
    if not book.refusals_complete:
        return "the refused sheets in hand do not match their census"
    if book.invariant:
        return "; ".join(dict.fromkeys(book.invariant.values()))
    if book.questions and book.reading is None:
        return ("a drawing-set question is open (which plans belong to the set): "
                + ", ".join(sorted(book.questions.values())))
    if book.opening and any(S._unread(r) for r in fam.rows):
        return "a tag or symbol of this schedule could not be read"
    if not n:
        q = fam.qty.get(tags[0]) if len(tags) == 1 else None
        return ("nothing located" + (f"; the schedule prints QTY {q}" if q else ""))
    st, q = S._qty_check(fam, tags, n)
    if st == "disagrees":
        return (f"{n} located against the schedule's QTY {q}: they disagree, "
                "so neither is stated")
    return "nothing located"


def _check_text(fam: S.GlyphFamily, tags: Sequence[str], n: int,
                readings: Optional[Dict[str, List[str]]] = None) -> str:
    st, q = S._qty_check(fam, tags, n)
    if st == "unreadable":
        # A QTY CELL TWO READINGS DISAGREE ON (ruling 73): the total stands,
        # unchecked - and the line says so, with what the two readings said
        # (the schedule record's contested_cells, carried by the census).
        said = [r for t in tags for r in (readings or {}).get(t) or []]
        return ("not checked: schedule QTY unreadable"
                + (f" ({' or '.join(dict.fromkeys(said))})" if said else ""))
    return {"agrees": f"schedule QTY {q} agrees",
            "no_column": "no schedule QTY"}.get(st, f"schedule QTY {q}")


def _qty_readings(records: Sequence[Dict[str, Any]]) -> Dict[str, Dict[str, List[str]]]:
    """family -> tag -> the readings of its contested QTY cell, from the
    censuses (plan_search.glyph_census qty_readings)."""
    out: Dict[str, Dict[str, List[str]]] = {}
    for c in records:
        if c.get("record_type") == S.GLYPH_CENSUS and isinstance(c.get("qty_readings"), dict):
            for t, rs in c["qty_readings"].items():
                out.setdefault(str(c.get("family")), {}).setdefault(t, [])
                out[str(c.get("family"))][t] += [str(r) for r in rs or []
                                                 if str(r) not in out[str(c.get("family"))][t]]
    return out


# ── where a mechanical symbol was placed ───────────────────────────────────

_NOT_PLACED = "refused"
_OUTSIDE = "placed-non-unit"
#: payload.units.method of a sheet counted at sheet scope
#: (plan_glyph_pass.SHEET_SCOPE_UNITS)
_SHEET_SCOPE = "sheet_scope"


def _unplaced_sheet(r: Dict[str, Any]) -> Optional[str]:
    """The architectural sheet of a row counted at sheet scope because that
    partner names no unit (plan_glyph_pass.sheet_scope, SHEET_SCOPE_UNITS:
    no unit tags read) - WH-1 on the owner's A-105.00 - else None.

    A WHOLE-SHEET COUNT BINDS AT SHEET SCOPE (operator ruling 2026-10-07,
    restated as rulings 101/102, 2026-10-28, withdrawing 72's refusal): in
    the export as in the chat, the floor line and the set's total stand, and
    the floor line says the symbols were not placed in apartments. A sheet
    with no architectural partner (SAF-1's M-106.00) and a plan-only family's
    plan are not this, and neither is one symbol a placed sheet could not
    place ("in a wall")."""
    p = r.get("payload") if isinstance(r.get("payload"), dict) else {}
    a = p.get("arch") if isinstance(p.get("arch"), dict) else None
    u = p.get("units") if isinstance(p.get("units"), dict) else {}
    if (a and a.get("sheet_number") and u.get("method") == _SHEET_SCOPE
            and p.get("placement") == _NOT_PLACED and not u.get("tags")):
        return str(a["sheet_number"])
    return None


def _placement_note(rows: Sequence[Dict[str, Any]]) -> str:
    """Where a floor line's symbols were placed, when all of them were placed
    the same way: not in apartments (a sheet-scope count), or outside them."""
    sheets = [_unplaced_sheet(r) for r in rows]
    if rows and all(sheets):
        return (f"not placed in apartments: {', '.join(dict.fromkeys(sheets))} "
                f"has no unit tags")
    return "placed outside apartments" if _outside_apartments(rows) else ""


def _outside_apartments(rows: Sequence[Dict[str, Any]]) -> bool:
    return bool(rows) and all((r.get("payload") or {}).get("placement") == _OUTSIDE for r in rows)


# ── the books ──────────────────────────────────────────────────────────────

class _Ctx(NamedTuple):
    records: List[Dict[str, Any]]
    sources: Dict[str, dict]
    src_list: List[dict]
    out: List[Line]
    refused: List[Line]
    readings: Dict[str, Dict[str, List[str]]] = {}


def _put(ctx: _Ctx, *, trade, item, tag, scope, level, n, unit, set_label, rows,
         check, sentence, why, intent="count") -> None:
    sheets, src = _source(rows, ctx.sources)
    if why is None:
        ok, bad = S.answer_is_grounded(sentence, ctx.records, intent=intent,
                                       sources=ctx.src_list)
        if ok:
            ctx.out.append(Line(trade, item, tag, scope, level, n, unit, set_label,
                                sheets, src, check, sentence))
            return
        why = f"the plan check did not bind {', '.join(bad) or 'this figure'}"
    ctx.refused.append(Line(trade, item, tag, scope, level, None, unit, set_label,
                            sheets, src, check, sentence if n else "", why))


def _book_lines(ctx: _Ctx, book: S.GlyphBook, *, set_label: str, sets: int,
                trade_of, per_set: bool = False, skip=None) -> None:
    """`per_set`: `book` is one drawing set's (glyph_books().tracks), and
    every line is per that set and says so, as a door or window line is.
    `skip(fam)`: a family this set gives no line for."""
    opening = book.opening
    refused_k = S._refused_floors(book)
    for name, fam in sorted(book.families.items()):
        if skip is not None and skip(fam):
            continue
        trade = trade_of(fam)
        phrase = _set_phrase(set_label, sets) if (opening or per_set) else ""
        # WHERE A MECHANICAL SYMBOL WAS PLACED is said on its floor line
        # (rulings 72, 101/102): outside the apartments, or not in them.
        mech = not opening and fam.plan_only is None

        def label_of(rows):
            # A MECHANICAL SYMBOL IS PER THE SET OF THE PLAN IT IS PLACED ON:
            # rows written before EMIT_VERSION 7 carry no track, and are placed
            # against the filed set on one floor and the owner's on another.
            if opening or per_set:
                return set_label
            return "; ".join(sorted({_row_set(r, ctx.sources) for r in rows})) or set_label
        # A FLOOR NO COUNT IS READ ON: one line for the family, not one per tag.
        fam_floors = {S._row_floor(r) for r in fam.rows if S._row_floor(r)}
        for k in sorted(refused_k & (fam_floors | set(book.titles))):
            why = _floor_why(book, fam, k)
            _put(ctx, trade=trade, item=name, tag="(all)", scope=FLOOR,
                 level=book.titles.get(k) or k, n=None, unit="ea", set_label=set_label,
                 rows=[r for r in book.refused if S._row_floor(r) == k], check="",
                 sentence="", why=why)
        if not fam.tags:
            # A SCHEDULE WITH NO IDENTIFIER COLUMN counts nothing - and is
            # never left out silently (plan_set_health.unreconciled).
            _put(ctx, trade=trade, item=name, tag="(all)", scope=BUILDING, level="", n=None,
                 unit="ea", set_label=set_label, rows=fam.rows, check="", sentence="",
                 why=_UNRECONCILED)
            continue
        for tag in fam.tags:
            rows_t = [r for r in fam.rows if r.get("label") == tag]
            for k in sorted({S._row_floor(r) for r in rows_t if S._row_floor(r)} - refused_k):
                on_k = [r for r in rows_t if S._row_floor(r) == k]
                n = S._floor_count(book, fam, [tag], k)
                subject = _floor_subject(book, k)
                why = None if n and subject else (
                    _floor_why(book, fam, k) if not n else "the floor has no name a check can read")
                _put(ctx, trade=trade, item=name, tag=S._type_name(tag) if opening else tag,
                     scope=FLOOR, level=book.titles.get(k) or k, n=n, unit="ea",
                     set_label=label_of(on_k), rows=on_k,
                     check=_placement_note(on_k) if mech else "",
                     sentence=_sentence(subject, n, _thing(fam, tag, n, opening), phrase)
                     if n else "", why=why)
            n = S._building_count(book, fam, [tag])
            located = len(S._resolved(fam, [tag]))
            if not n and not located and not fam.qty.get(tag) and not fam.why \
                    and not S._building_refused(book):
                continue                # a schedule type nothing names: no line at all
            why = None if n else _building_why(book, fam, [tag], located)
            _put(ctx, trade=trade, item=name, tag=S._type_name(tag) if opening else tag,
                 scope=BUILDING, level="", n=n, unit="ea", set_label=label_of(rows_t),
                 rows=rows_t,
                 check=_check_text(fam, [tag], located, ctx.readings.get(name)) if located else "",
                 sentence=_sentence("", n, _thing(fam, tag, n, opening), phrase) if n else "",
                 why=why)


def _not_read_by(records: Sequence[Dict[str, Any]], fam: S.GlyphFamily, t: Any,
                 books: S.GlyphBooks) -> bool:
    """A family this set's pass located nothing of because it does not read
    the sheets the other sets located it on (SAF-1 on M-106.00, the bulkhead,
    read by the filed set alone): no line for it in this set - the same test
    as plan_search._unread_by_other_tracks. A set that reads the sheet and
    locates none still gets its line, refused."""
    if fam.rows:
        return False
    cs = [c for c in records if c.get("record_type") == S.GLYPH_CENSUS
          and c.get("kind") == "family" and c.get("family") == fam.name and c.get("track") == t]
    if len(cs) != 1 or cs[0].get("count") != 0 or not isinstance(cs[0].get("reads"), dict):
        return False
    reads = cs[0]["reads"]
    pages, sheets = set(reads.get("pages") or ()), set(reads.get("sheets") or ())
    elsewhere = [r for o, b in books.tracks.items() if o != t
                 for r in (b.families[fam.name].rows if fam.name in b.families else [])]
    return bool(elsewhere) and not any(
        str(r.get("page_id")) in pages or not r.get("sheet_number")
        or str(r.get("sheet_number")) in sheets for r in elsewhere)


_UNRECONCILED = "the schedule has no identifier column: its items cannot be counted"


def _unreconciled_lines(ctx: _Ctx, records: Sequence[Dict[str, Any]]) -> None:
    """Every schedule set health says has no identifier column, that no book
    already listed: one refused line each, with the recorded reason."""
    from lib import plan_set_health as _health
    done = {ln.item for ln in ctx.refused}
    for h in records:
        if h.get("record_type") != "set_health" or h.get("kind") != _health.FAMILY_NOT_RECONCILED:
            continue
        name = str(h.get("subject") or "").strip()
        if not name or name in done:
            continue
        done.add(name)
        trade = (DOORS_WINDOWS if re.search(r"DOOR|WINDOW", name, re.I)
                 else _trade(S.GlyphFamily(name, [], [], None)))
        sheet = str(h.get("sheet") or "")
        ctx.refused.append(Line(trade, name, "(all)", BUILDING, "", None, "ea", "", sheet,
                                sheet, "", "", str(h.get("detail") or _UNRECONCILED)))


def _area_lines(ctx: _Ctx, rows: Sequence[Dict[str, Any]]) -> None:
    rows = [r for r in rows if _areas.is_area_record(r)]
    tracks = {_areas._row_track(r) for r in rows}
    for r in sorted(rows, key=lambda r: (str(r.get("unit") or "~"),
                                         (r.get("payload") or {}).get("scope") != _areas.APARTMENT,
                                         str((r.get("payload") or {}).get("name") or ""),
                                         str(r.get("sheet_number") or ""))):
        p = _areas._payload(r)
        apt = p.get("scope") == _areas.APARTMENT
        unit, name = r.get("unit"), p.get("name") or ""
        label = p.get("drawing_track_label") or (
            S._FILED_LABEL if p.get("track") in (None, S._FILED) else f"the {p.get('track')} set")
        printed = _areas._figure(p.get("printed_sf"))
        who = (f"Apartment {unit}" if apt and unit else
               f"The {name.lower()} in {unit}" if unit else f"The {name.lower()}")
        st = r.get("status")
        phrase = _set_phrase(label, len(tracks))
        sentence, why, check = "", None, ""
        if st == _areas.AGREES:
            sentence = f"{who} is {printed} sq ft{phrase}."
            check = (f"measured {p.get('computed_sf')} sq ft at "
                     f"{(p.get('viewport') or {}).get('note') or 'its verified scale'} "
                     f"({p.get('diff_pct'):+.1f}%)")
        elif st == _areas.DISAGREES:
            why = (f"printed {printed} sq ft, but the drawn region measures "
                   f"{p.get('computed_sf')} sq ft ({p.get('diff_pct'):+.1f}%): neither is stated")
        elif st == _areas.NO_PRINTED:
            why = "no area is printed for it on this sheet"
        else:
            why = f"not checked: {p.get('why') or 'no area can be measured'}"
        _put(ctx, trade=AREAS, item=name if not apt else "Apartment NET",
             tag=str(unit or name), scope="Apartment" if apt else "Room", level="",
             n=float(printed) if st == _areas.AGREES and printed else None, unit="sq ft",
             set_label=label, rows=[r], check=check, sentence=sentence, why=why,
             intent="attribute")


def _legend_lines(ctx: _Ctx, rows: Sequence[Dict[str, Any]]) -> None:
    """A STORED LEGEND SYMBOL COUNT IS NEVER A FIGURE (ruling 77). Each counted
    row (lib.plan_legend_symbols.legend_count_rows, a count above zero) is one
    "Not verified - refused" line - its sheet, file and page, and its set when
    the row says one - with LEGEND_NOT_VERIFIED as the reason, and so one
    review item. No quantity is shown and nothing is put to the gate: its
    basis, legend_symbol, is not a counting basis. A refused row or a zero is
    the pass's own record and makes no line."""
    from lib import plan_legend_symbols as _pls
    for r in sorted(rows, key=lambda r: (str(r.get("sheet_number") or ""), r.get("ordinal") or 0)):
        if (r.get("record_type") != _pls.RECORD_TYPE or r.get("status") != "counted"
                or not r.get("count")):
            continue
        sheets, src = _source([r], ctx.sources)
        fsrc = ctx.sources.get(str(r.get("file_id") or "")) or {}
        t = r.get("track")
        set_label = fsrc.get("set_name") or (f"the {t} set" if t else "")
        ctx.refused.append(Line(LEGEND_SYMBOLS, str(r.get("label") or ""),
                                str(r.get("mark") or ""), "Sheet", sheets, None, "ea",
                                set_label, sheets, src, "", "", LEGEND_NOT_VERIFIED))


def _light_fixture_lines(ctx: _Ctx, rows: Sequence[Dict[str, Any]]) -> None:
    """A STORED LIGHT FIXTURE TAG COUNT IS NEVER A FIGURE (T6, ruling 77 as
    applied to T2's legend counts). Each counted row of
    lib.plan_light_fixtures.fixture_rows (a count above zero) is one "Not
    verified - refused" line - its drawing's level, its sheet, file and page,
    its set - with plan_light_fixtures.NOT_VERIFIED as the reason, and so one
    review item; a conflict row (the count disagrees with a printed total) is
    a refused line with its own reason. No quantity is shown and nothing is
    put to the gate: its basis, fixture_tag, is not a counting basis. A
    refused tag, an untagged drawing or a zero is the pass's own record and
    makes no line."""
    from lib import plan_light_fixtures as _plf
    for r in sorted(rows, key=lambda r: (str(r.get("sheet_number") or ""), str(r.get("title") or ""),
                                         str(r.get("tag") or ""))):
        if r.get("record_type") != _plf.RECORD_TYPE or r.get("status") not in ("counted", "conflict"):
            continue
        if r.get("status") == "counted" and not r.get("count"):
            continue
        sheets, src = _source([r], ctx.sources)
        fsrc = ctx.sources.get(str(r.get("file_id") or "")) or {}
        t = r.get("track")
        set_label = fsrc.get("set_name") or (f"the {t} set" if t else "")
        level = " / ".join(r.get("level_names") or []) or str(r.get("title") or "")
        why = _plf.NOT_VERIFIED if r.get("status") == "counted" else str(r.get("refused") or "")
        ctx.refused.append(Line(_plf.TRADE, "Light fixture", str(r.get("tag") or ""),
                                FLOOR if r.get("levels") else "Sheet", level, None, "ea",
                                set_label, sheets, src, "", "", why))


def build(records: Sequence[Dict[str, Any]], area_rows: Sequence[Dict[str, Any]] = (),
          sources: Optional[Dict[str, dict]] = None,
          legend_rows: Sequence[Dict[str, Any]] = (),
          light_rows: Sequence[Dict[str, Any]] = ()) -> Takeoff:
    """Every candidate line, put to the gate. `records`: the located-symbol
    evidence (rows, censuses, set health) of every family; `area_rows`: the
    project's area rows; `sources`: {file_id: plan_sources.source_of};
    `legend_rows`: stored legend symbol counts, each listed as not verified
    (_legend_lines); `light_rows`: stored light fixture tag counts, likewise
    (_light_fixture_lines) - none stored, nothing changes."""
    records = list(records or [])
    sources = dict(sources or {})
    ctx = _Ctx(records, sources, list(sources.values()), [], [], _qty_readings(records))
    gb = S.glyph_books(records)
    if gb.totals is not None and len(gb.tracks) > 1:
        # ONE LINE PER SET (ruling 72), as the door and window lines are.
        for t, book in gb.tracks.items():
            _book_lines(ctx, book, set_label=gb.labels.get(t) or S._FILED_LABEL,
                        sets=len(gb.tracks), trade_of=_trade, per_set=True,
                        skip=lambda f, t=t: _not_read_by(records, f, t, gb))
    elif gb.totals is not None:
        labels = gb.labels
        one = "; ".join(labels.values()) or S._FILED_LABEL
        _book_lines(ctx, gb.totals, set_label=one, sets=len(labels), trade_of=_trade)
    ob = S.opening_books(records)
    for t, book in ob.tracks.items():
        _book_lines(ctx, book, set_label=ob.labels.get(t) or S._FILED_LABEL,
                    sets=len(ob.tracks), trade_of=lambda f: DOORS_WINDOWS)
    _unreconciled_lines(ctx, records)
    if area_rows:
        actx = _Ctx(list(area_rows), sources, ctx.src_list, ctx.out, ctx.refused)
        _area_lines(actx, area_rows)
    if legend_rows:
        _legend_lines(ctx, legend_rows)
    if light_rows:
        _light_fixture_lines(ctx, light_rows)
    return Takeoff(ctx.out, ctx.refused)


# ── the workbook ───────────────────────────────────────────────────────────

_COLUMNS = ("Item", "Tag / type", "Scope", "Level", "Qty", "Unit", "Set", "Sheets",
            "Source (file, page)", "Cross-check", "Verified as")
_REFUSED_COLUMNS = ("Trade", "Item", "Tag / type", "Scope", "Level", "Set", "Sheets",
                    "Source (file, page)", "Reason")


def write(takeoff: Takeoff, *, project_name: str = "", generated: str = "",
          code_sha: str = "") -> bytes:
    """The .xlsx: one sheet per trade, then the refused sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    bold = Font(bold=True)
    head = (f"Takeoff - {project_name}" if project_name else "Takeoff")
    note = ("Every figure here is one the plan check binds to the drawings; anything it "
            "could not verify is on the '" + REFUSED_SHEET + "' sheet with its reason.")
    meta = " · ".join(x for x in (f"Generated {generated}" if generated else "",
                                   f"build {code_sha[:12]}" if code_sha else "") if x)

    def sheet(ws, columns, rows):
        ws.append([head])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([note])
        ws.append([meta])
        ws.append(list(columns))
        for c in range(1, len(columns) + 1):
            ws.cell(row=4, column=c).font = bold
        for r in rows:
            ws.append(list(r))
        if not rows:
            ws.append(["None found."])
        ws.freeze_panes = "A5"
        for i, col in enumerate(columns, 1):
            width = max([len(str(col))] + [len(str(r[i - 1] or "")) for r in rows] or [10])
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)

    first = True
    trades = takeoff.by_trade()
    for trade in list(TRADES) + [t for t in trades if t not in TRADES]:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = trade[:31]
        sheet(ws, _COLUMNS, [
            (ln.item, ln.tag, ln.scope, ln.level,
             int(ln.qty) if ln.qty is not None and float(ln.qty).is_integer() else ln.qty,
             ln.unit, ln.set_label, ln.sheets, ln.source, ln.check, ln.sentence)
            for ln in trades.get(trade, [])])
    ws = wb.create_sheet(REFUSED_SHEET[:31])
    sheet(ws, _REFUSED_COLUMNS, [
        (ln.trade, ln.item, ln.tag, ln.scope, ln.level, ln.set_label, ln.sheets,
         ln.source, ln.reason) for ln in takeoff.refused])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
