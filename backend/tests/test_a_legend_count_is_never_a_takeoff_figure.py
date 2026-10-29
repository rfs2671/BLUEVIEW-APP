"""A stored legend symbol count is never a takeoff figure (ruling 77).

Roy, 2026-10-28: when legend_count rows exist, the takeoff export (T5,
lib/plan_takeoff_export) lists each as a "Not verified - refused" line,
reason "legend symbol count: not verified until the sub-bid benchmark (drawn
symbols, sheet scope)", and queues a review item through
review_queue.from_takeoff (#911). It is never a verified figure, and
`legend_symbol` stays out of plan_search.COUNTING_BASES. With no rows stored
- no pass has run (ruling 78) - the export is exactly what it was.

The world is T5's own (tests/test_a_takeoff_line_is_one_the_gate_binds.py).
"""
from __future__ import annotations

import io
import os
import sys
import unittest
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "smoke_test")
os.environ.setdefault("JWT_SECRET", "smoke_test_secret")
os.environ.setdefault("APP_BASE_URL", "https://app.levelog.com")

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import server  # noqa: E402
from lib import plan_search as S  # noqa: E402
from lib import plan_takeoff_export as X  # noqa: E402
from lib import review_queue  # noqa: E402
from tests import test_a_takeoff_line_is_one_the_gate_binds as T5  # noqa: E402

REASON = ("legend symbol count: not verified until the sub-bid benchmark "
          "(drawn symbols, sheet scope)")


def _need(mod, *names):
    for n in names:
        assert hasattr(mod, n), f"{mod.__name__} has no {n}"


def _rows():
    """What lib.plan_legend_symbols.legend_count_rows stores for P-202.00:
    three counted stacks, a zero, and a refused swatch."""
    base = {"project_id": "proj1", "file_id": "f-pl", "page_number": 10,
            "sheet_number": "P-202.00", "file_name": "PL - 6.29.26.pdf",
            "record_type": "legend_symbol_count", "count_basis": "legend_symbol",
            "scope": "sheet", "track": None}
    return [dict(base, ordinal=0, label="SANITARY STACK/STACK NO.", status="counted", count=4),
            dict(base, ordinal=1, label="VENT STACK/STACK NO.", status="counted", count=4),
            dict(base, ordinal=2, label="STORM STACK/STACK NO.", status="counted", count=2),
            dict(base, ordinal=3, label="HOSE BIBB", status="counted", count=0),
            dict(base, ordinal=4, label="SANITARY PIPING", status="refused", count=None,
                 refused="a line type (a run drawn across the cell), not a placed symbol")]


def _build(legend=None):
    if legend is None:
        return X.build(T5._world(), T5.AREAS, T5.SOURCES)
    return X.build(T5._world(), T5.AREAS, T5.SOURCES, legend_rows=legend)


class EachCountIsANotVerifiedLine(unittest.TestCase):

    def test_each_counted_row_is_one_refused_line_with_the_ruled_reason(self):
        _need(X, "LEGEND_NOT_VERIFIED", "LEGEND_SYMBOLS")
        self.assertEqual(X.LEGEND_NOT_VERIFIED, REASON)
        t = _build(_rows())
        mine = [ln for ln in t.refused if ln.trade == X.LEGEND_SYMBOLS]
        self.assertEqual(sorted(ln.item for ln in mine),
                         ["SANITARY STACK/STACK NO.", "STORM STACK/STACK NO.",
                          "VENT STACK/STACK NO."])
        for ln in mine:
            self.assertEqual(ln.reason, REASON)
            self.assertIsNone(ln.qty)                  # never a figure
            self.assertEqual((ln.scope, ln.level, ln.sheets), ("Sheet", "P-202.00", "P-202.00"))
            self.assertIn("PL - 6.29.26.pdf p10", ln.source)
            self.assertEqual(ln.sentence, "")           # nothing was put to the gate
            self.assertEqual(ln.set_label, "")          # no set claimed for an untracked row

    def test_never_a_verified_figure(self):
        _need(X, "LEGEND_SYMBOLS")
        t = _build(_rows())
        self.assertFalse([ln for ln in t.lines if ln.trade == X.LEGEND_SYMBOLS])
        self.assertNotIn("legend_symbol", S.COUNTING_BASES)

    def test_a_zero_and_a_refused_swatch_make_no_line(self):
        _need(X, "LEGEND_SYMBOLS")
        items = {ln.item for ln in _build(_rows()).refused if ln.trade == X.LEGEND_SYMBOLS}
        self.assertNotIn("HOSE BIBB", items)
        self.assertNotIn("SANITARY PIPING", items)

    def test_each_line_makes_a_review_item(self):
        _need(X, "LEGEND_SYMBOLS")
        ln = [ln for ln in _build(_rows()).refused if ln.trade == X.LEGEND_SYMBOLS][0]
        doc = review_queue.from_takeoff("co_a", "proj1", trade=ln.trade,
                                        item_name=f"{ln.item} {ln.tag}".strip(),
                                        scope=" ".join(x for x in (ln.scope, ln.level) if x),
                                        track=ln.set_label, reason=ln.reason,
                                        sheets=[ln.sheets], now=datetime.now(timezone.utc))
        self.assertIsNotNone(doc)
        self.assertEqual(doc["reason"], review_queue.REASON_TAKEOFF_REFUSED)


class WithNothingStoredTheExportIsUnchanged(unittest.TestCase):

    def test_no_rows_builds_the_same_takeoff(self):
        _need(X, "LEGEND_SYMBOLS")
        self.assertEqual(_build(), _build([]))

    def test_the_route_loads_them_and_with_none_writes_the_same_lines(self):
        _need(server, "_legend_count_evidence")
        r0, q0 = _route(())
        r1, q1 = _route(_rows())
        self.assertEqual((r0.status_code, r1.status_code), (200, 200))
        mine = [line for _c, _p, line in q1 if line["reason"] == REASON]
        self.assertEqual(len(mine), 3)
        self.assertTrue(all(line["trade"] == "Legend symbols" for line in mine))
        # nothing stored: the queued lines and the refused sheet are what they were
        self.assertFalse([line for _c, _p, line in q0 if line["reason"] == REASON])
        self.assertEqual([line for _c, _p, line in q1 if line["reason"] != REASON],
                         [line for _c, _p, line in q0])
        ws0 = load_workbook(io.BytesIO(r0.content))[X.REFUSED_SHEET[:31]]
        ws1 = load_workbook(io.BytesIO(r1.content))[X.REFUSED_SHEET[:31]]
        self.assertEqual(ws1.max_row - ws0.max_row, 3)


def _route(legend):
    async def _user():
        return {"_id": "u1", "id": "u1", "role": "admin", "company_id": "co_a",
                "account_status": "approved", "full_name": "Test", "assigned_projects": []}
    queued = []

    async def _queue(company_id, project_id, **line):
        queued.append((company_id, project_id, line))
        return True

    async def _glyphs(pid, subject, *, everything=False):
        return T5._world()

    async def _areas(pid):
        return T5.AREAS

    async def _sources(pid, fids):
        return T5.SOURCES

    async def _legend(pid):
        return list(legend)

    async def _lights(pid):
        # stored light fixture tag counts (T6): none here
        return []

    server.app.dependency_overrides[server.get_current_user] = _user
    try:
        with patch.object(server, "db", T5._Db(T5.PROJECT)), \
                patch.object(server, "_glyph_evidence", _glyphs), \
                patch.object(server, "_area_evidence", _areas), \
                patch.object(server, "_plan_file_sources", _sources), \
                patch.object(server, "_legend_count_evidence", _legend), \
                patch.object(server, "_light_fixture_evidence", _lights), \
                patch.object(server, "_review_queue_takeoff", _queue):
            return TestClient(server.app).get("/api/projects/proj1/takeoff.xlsx"), queued
    finally:
        server.app.dependency_overrides.clear()


if __name__ == "__main__":
    unittest.main()
