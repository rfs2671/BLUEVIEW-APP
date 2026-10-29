"""T5, THE TAKEOFF EXPORT: every figure in the workbook is one the plan gate
binds, every line carries its sheets, file and page, and set, and every
figure the gate does not bind is on the refused sheet with its reason.

The worlds are the gate's own fixtures: two drawing tracks of PTAC rows
(tests/test_two_drawing_tracks.py), door tags in two sets
(tests/test_a_door_is_counted_by_its_tag.py), and area rows
(tests/test_an_area_binds_only_where_the_plan_prints_one_that_agrees.py).
"""

from __future__ import annotations

import io
import os
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "smoke_test")
os.environ.setdefault("JWT_SECRET", "smoke_test_secret")
os.environ.setdefault("APP_BASE_URL", "https://app.levelog.com")

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import server  # noqa: E402
from lib import plan_search as S  # noqa: E402
from lib import plan_takeoff_export as X  # noqa: E402
from tests import test_a_door_is_counted_by_its_tag as D  # noqa: E402
from tests import test_an_area_binds_only_where_the_plan_prints_one_that_agrees as A  # noqa: E402
from tests import test_two_drawing_tracks as T  # noqa: E402

OWNER = T.LABEL
AREAS = [r for r in A.WORLD if r.get("record_type") == "area"]
SOURCES = {"f-mh": {"file_id": "f-mh", "file_name": "MH.pdf", "set_name": "filed set"},
           "f-filed": {"file_id": "f-filed", "file_name": "AR - 8.18.26.pdf",
                       "set_name": "filed set"},
           "f-post_co": {"file_id": "f-post_co", "file_name": "Owners set - 6.9.26.pdf",
                         "set_name": "owner's set"}}


def _world():
    return T._world() + D._world(D._first_floor("filed"), D._first_floor("post_co"))


def _find(lines, *, tag, scope, set_has=None, level=None):
    return [ln for ln in lines if ln.tag == tag and ln.scope == scope
            and (set_has is None or set_has in ln.set_label)
            and (level is None or ln.level == level)]


class EveryExportedFigureIsBound(unittest.TestCase):

    def test_each_line_binds_through_the_gate_and_states_its_figure(self):
        recs = _world()
        t = X.build(recs, AREAS)
        self.assertTrue(t.lines)
        for ln in t.lines:
            intent = "attribute" if ln.trade == X.AREAS else "count"
            ok, bad = S.answer_is_grounded(ln.sentence, recs if intent == "count" else AREAS,
                                           intent=intent)
            self.assertTrue(ok, (ln.sentence, bad))
            q = int(ln.qty) if float(ln.qty).is_integer() else ln.qty
            self.assertIn(f" {q} ", f" {ln.sentence} ", ln)

    def test_every_line_has_a_source_and_a_set(self):
        t = X.build(_world(), AREAS, SOURCES)
        for ln in t.lines:
            self.assertTrue(ln.sheets and ln.source and ln.set_label, ln)
        ptac = _find(t.lines, tag="PTAC-1", scope=X.BUILDING)[0]
        self.assertEqual(ptac.qty, 3)
        self.assertIn("M-100.00: MH.pdf p5", ptac.source)
        self.assertIn("A-100.00: AR - 8.18.26.pdf p4", ptac.source)

    def test_the_figures_the_gate_states(self):
        """One line per drawing set (ruling 72), as a door or window line is."""
        t = X.build(_world())
        self.assertEqual({ln.set_label: ln.qty
                          for ln in _find(t.lines, tag="PTAC-3", scope=X.BUILDING)},
                         {"the filed set": 3, OWNER: 3})
        self.assertEqual({ln.set_label: ln.qty
                          for ln in _find(t.lines, tag="PTAC-1", scope=X.FLOOR,
                                          level="SECOND FLOOR HVAC PLANS")},
                         {"the filed set": 2, OWNER: 2})
        self.assertEqual(sum(ln.trade == X.EQUIPMENT for ln in t.lines), 16)

    def test_a_door_type_is_per_its_set(self):
        """Type F is drawn in the owner's set only: one line, per that set,
        and its sentence names the set - as the gate requires."""
        t = X.build(_world())
        f = _find(t.lines, tag="type F", scope=X.BUILDING)
        self.assertEqual([(ln.qty, ln.set_label) for ln in f], [(1, OWNER)])
        self.assertIn(OWNER, f[0].sentence)
        self.assertEqual(f[0].trade, X.DOORS_WINDOWS)
        g = {ln.set_label: ln.qty for ln in _find(t.lines, tag="type G", scope=X.BUILDING)}
        self.assertEqual(g, {"the filed set": 3, OWNER: 4})

    def test_areas_only_where_printed_and_measured_agree(self):
        t = X.build([], AREAS)
        self.assertEqual(sorted(ln.tag for ln in t.lines), ["2A", "2C", "BIKE ROOM"])
        self.assertEqual(_find(t.lines, tag="2A", scope="Apartment")[0].qty, 482)
        why = {ln.tag: ln.reason for ln in t.refused}
        self.assertIn("printed 338 sq ft", why["1A"])
        self.assertIn("451.0 sq ft", why["1A"])
        self.assertIn("no area is printed", why["4A"])
        self.assertIn("not verified", why["4B"])
        self.assertNotIn(338, [ln.qty for ln in t.lines])


class WhatIsNotBoundIsRefusedWithItsReason(unittest.TestCase):

    def test_a_floor_whose_track_is_short_of_its_sheet(self):
        with self.assertLogs("lib.plan_search", "WARNING"):
            t = X.build(T._world(owner_100=T.OWNER_100[:4], owner_mech=5))
        # the owner's set is short on its first floor: ITS lines there and
        # its totals are refused ...
        self.assertFalse(_find(t.lines, tag="PTAC-1", scope=X.BUILDING, set_has="owner"))
        self.assertFalse([ln for ln in t.lines if ln.level == "FIRST FLOOR HVAC PLANS"
                          and ln.set_label == OWNER])
        floor = [ln for ln in t.refused if ln.level == "FIRST FLOOR HVAC PLANS"]
        self.assertEqual([ln.set_label for ln in floor], [OWNER])
        self.assertIn("4 located on the post_co track against 5", floor[0].reason)
        for tag in T.TAGS:
            self.assertIn("against 5", _find(t.refused, tag=tag, scope=X.BUILDING)[0].reason)
        # ... and the filed set, counted on its own (ruling 72), stands
        self.assertEqual([(ln.qty, ln.set_label)
                          for ln in _find(t.lines, tag="PTAC-1", scope=X.BUILDING)],
                         [(3, "the filed set")])
        self.assertEqual(sorted((ln.set_label, ln.level, ln.qty)
                                for ln in _find(t.lines, tag="PTAC-1", scope=X.FLOOR)),
                         [("the filed set", "FIRST FLOOR HVAC PLANS", 1),
                          ("the filed set", "SECOND FLOOR HVAC PLANS", 2),
                          (OWNER, "SECOND FLOOR HVAC PLANS", 2)])

    def test_a_total_the_schedule_qty_contradicts(self):
        t = X.build(T._untracked(qty={"PTAC-1": "9"}))
        self.assertFalse(_find(t.lines, tag="PTAC-1", scope=X.BUILDING))
        why = _find(t.refused, tag="PTAC-1", scope=X.BUILDING)[0].reason
        self.assertIn("3 located against the schedule's QTY 9", why)

    def test_a_schedule_with_nothing_located(self):
        t = X.build(_world())
        w = [ln for ln in t.refused if ln.item == "WINDOW SCHEDULE"]
        self.assertEqual(len(w), 6)               # W1-W3, in each set
        self.assertTrue(all(ln.reason == "no located symbols" for ln in w))

    def test_a_schedule_with_no_identifier_column_is_listed(self):
        """Codex, #912: an unreconciled family counted nothing and was left
        out of the workbook - "None found." where the set says otherwise."""
        from lib import plan_set_health as H
        health = [dict(h, record_type="set_health") for h in
                  H.unreconciled({"DIFFUSER SCHEDULE": {"sheet": "M-001.00", "tags": []}})]
        census = S.glyph_census("family", T.PID, ["pg-m100"], 0, family="DIFFUSER SCHEDULE",
                                tags=[])
        for recs in (health, health + [census], T._world() + health):
            t = X.build(recs)
            d = [ln for ln in t.refused if ln.item == "DIFFUSER SCHEDULE"]
            self.assertEqual(len(d), 1, recs)
            self.assertEqual(d[0].trade, X.EQUIPMENT)
            self.assertIn("no identifier column", d[0].reason)
        self.assertEqual(d[0].sheets, "M-001.00")

    def test_a_figure_the_gate_refuses_is_never_exported(self):
        with patch.object(S, "answer_is_grounded", return_value=(False, ["3"])):
            t = X.build(_world(), AREAS)
        self.assertEqual(t.lines, [])
        self.assertTrue(all(ln.qty is None for ln in t.refused))
        self.assertIn("the plan check did not bind 3",
                      _find(t.refused, tag="PTAC-1", scope=X.BUILDING)[0].reason)

    def test_refused_lines_carry_their_source(self):
        t = X.build(T._world(owner_100=T.OWNER_100[:4], owner_mech=5), sources=SOURCES)
        b = _find(t.refused, tag="PTAC-1", scope=X.BUILDING)[0]
        self.assertIn("M-100.00", b.sheets)
        self.assertIn("MH.pdf", b.source)


class TheTrades(unittest.TestCase):

    def _fam(self, name, tags):
        return S.GlyphFamily(name, tags, [], None)

    def test_roof_drains(self):
        self.assertEqual(X._trade(self._fam("ROOF DRAIN SCHEDULE", ["RD-1"])), X.ROOF_DRAINS)
        self.assertEqual(X._trade(self._fam("PLUMBING", ["RD-1", "RD-2"])), X.ROOF_DRAINS)
        self.assertEqual(X._trade(self._fam("ROOMS PTAC UNITS SCHEDULE", T.TAGS)), X.EQUIPMENT)
        self.assertEqual(X._trade(self._fam("FAN SCHEDULE", ["RD-1", "EF-1"])), X.EQUIPMENT)

    def test_a_floor_is_named_by_its_level_else_its_sheet(self):
        b = S.glyph_books(T._world()).totals
        self.assertEqual(X._floor_subject(b, "100"), "the first floor")
        b2 = b._replace(levels={"100": {1}, "101": {1}})
        self.assertEqual(X._floor_subject(b2, "100"), "sheet M-100.00")


class TheWorkbook(unittest.TestCase):

    def test_one_sheet_per_trade_then_the_refused(self):
        t = X.build(_world(), AREAS, SOURCES)
        wb = load_workbook(io.BytesIO(X.write(t, project_name="Fixture", code_sha="abc123")))
        self.assertEqual(wb.sheetnames, list(X.TRADES) + [X.REFUSED_SHEET])
        eq = wb[X.EQUIPMENT]
        self.assertEqual([c.value for c in eq[4]], list(X._COLUMNS))
        rows = [[c.value for c in r] for r in eq.iter_rows(min_row=5)]
        self.assertIn(["ROOMS PTAC UNITS SCHEDULE", "PTAC-1", "Building", None, 3, "ea"],
                      [r[:6] for r in rows])
        self.assertEqual([[c.value for c in r] for r in wb[X.ROOF_DRAINS].iter_rows(min_row=5)],
                         [["None found."] + [None] * (len(X._COLUMNS) - 1)])
        ref = [[c.value for c in r] for r in wb[X.REFUSED_SHEET].iter_rows(min_row=5)]
        self.assertTrue(any(r[1] == "WINDOW SCHEDULE" and r[-1] == "no located symbols"
                            for r in ref))
        self.assertIn("abc123", wb[X.EQUIPMENT]["A3"].value)


# ── the route ──────────────────────────────────────────────────────────────

class _Col:
    def __init__(self, doc=None):
        self.doc = doc

    async def find_one(self, *a, **k):
        return self.doc


class _Db:
    def __init__(self, project):
        self.projects = _Col(project)

    def __getattr__(self, n):
        if n.startswith("_"):
            raise AttributeError(n)
        return _Col()

    def __getitem__(self, n):
        return _Col()


PROJECT = {"_id": "proj1", "company_id": "co_a", "name": "588 Boyland"}


def _get(role, company="co_a", legend=(), lights=()):
    async def _user():
        return {"_id": "u1", "id": "u1", "role": role, "company_id": company,
                "account_status": "approved", "full_name": "Test", "assigned_projects": []}

    queued = []

    async def _queue(company_id, project_id, **line):
        queued.append((company_id, project_id, line))
        return True

    async def _glyphs(pid, subject, *, everything=False):
        assert everything
        return _world()

    async def _areas(pid):
        return AREAS

    async def _sources(pid, fids):
        return SOURCES

    async def _legend(pid):
        # stored legend symbol counts (lib.plan_legend_symbols): none until a
        # pass runs (ruling 78)
        return list(legend)

    async def _lights(pid):
        # stored light fixture tag counts (lib.plan_light_fixtures): none
        # until a pass runs (ruling 78)
        return list(lights)

    server.app.dependency_overrides[server.get_current_user] = _user
    try:
        with patch.object(server, "db", _Db(PROJECT)), \
                patch.object(server, "_glyph_evidence", _glyphs), \
                patch.object(server, "_area_evidence", _areas), \
                patch.object(server, "_plan_file_sources", _sources), \
                patch.object(server, "_legend_count_evidence", _legend),                 patch.object(server, "_light_fixture_evidence", _lights),                 patch.object(server, "_review_queue_takeoff", _queue):
            return TestClient(server.app).get("/api/projects/proj1/takeoff.xlsx"), queued
    finally:
        server.app.dependency_overrides.clear()


class TheRoute(unittest.TestCase):

    def test_an_admin_gets_the_workbook_and_refused_lines_are_queued(self):
        r, queued = _get("admin")
        self.assertEqual(r.status_code, 200, r.text[:300])
        self.assertIn("spreadsheetml", r.headers["content-type"])
        self.assertIn('filename="Levelog_takeoff_588_Boyland_', r.headers["content-disposition"])
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames[-1], X.REFUSED_SHEET)
        self.assertTrue(queued)
        self.assertTrue(all(c == "co_a" and p == "proj1" and line["reason"]
                            for c, p, line in queued))

    def test_a_worker_is_refused(self):
        r, queued = _get("worker")
        self.assertEqual(r.status_code, 403)
        self.assertEqual(queued, [])

    def test_another_companys_admin_is_refused(self):
        r, queued = _get("admin", company="co_b")
        self.assertIn(r.status_code, (403, 404))
        self.assertEqual(queued, [])


if __name__ == "__main__":
    unittest.main()
