"""Schedule import: every format, and the refusals.

Fixtures are built here (inline XER / XML / CSV text, an openpyxl workbook in
memory, a PDF drawn by weasyprint). No network, no files.
"""

from __future__ import annotations

import io
import socket
from datetime import date, datetime

import pytest

from lib import schedule_import as S
from tests import _weasyprint
from lib.schedule_import import ScheduleImportError, parse_schedule


def _by_id(doc):
    return {a["id"]: a for a in doc["activities"]}


# ---------------------------------------------------------------- XER

_XER_FIELDS_PROJ = ["proj_id", "proj_short_name", "last_recalc_date", "plan_start_date"]
_XER_FIELDS_WBS = ["wbs_id", "proj_id", "parent_wbs_id", "wbs_short_name", "wbs_name", "seq_num"]
_XER_FIELDS_TASK = [
    "task_id", "proj_id", "wbs_id", "clndr_id", "task_code", "task_name", "task_type",
    "complete_pct_type", "phys_complete_pct", "target_drtn_hr_cnt",
    "target_start_date", "target_end_date", "early_start_date", "early_end_date",
    "act_start_date", "act_end_date",
]  # fmt: skip


def _xer(tasks, *, with_proj=True):
    def table(name, fields, rows):
        out = [f"%T\t{name}", "%F\t" + "\t".join(fields)]
        out += ["%R\t" + "\t".join(r) for r in rows]
        return "\n".join(out)

    parts = ["ERMHDR\t20.12\t2026-10-05\tProject\tadmin\tdb\tProject Management\tUSD"]
    if with_proj:
        parts.append(table("PROJECT", _XER_FIELDS_PROJ, [["1", "TWR", "2026-10-05 08:00", "2026-01-05 08:00"]]))
    parts.append(table("CALENDAR", ["clndr_id", "day_hr_cnt"], [["9", "8"]]))
    parts.append(table("PROJWBS", _XER_FIELDS_WBS, [
        ["10", "1", "", "TWR", "Tower 12", "1"],
        ["11", "1", "10", "SUB", "Substructure", "1"],
        ["12", "1", "10", "SUP", "Superstructure", "2"],
        ["13", "1", "12", "L3", "Level 3", "1"],
        ["14", "1", "10", "EMP", "Empty", "3"],
    ]))
    parts.append(table("TASK", _XER_FIELDS_TASK, tasks))
    parts.append("%E")
    return "\n".join(parts).encode("utf-8")


def _t(code, name, wbs, **kw):
    base = dict.fromkeys(_XER_FIELDS_TASK, "")
    base.update(task_id=code, proj_id="1", wbs_id=wbs, clndr_id="9", task_code=code, task_name=name,
                task_type="TT_Task", complete_pct_type="CP_Phys")  # fmt: skip
    base.update(kw)
    return [base[f] for f in _XER_FIELDS_TASK]


def test_xer_happy_path():
    data = _xer([
        _t("A100", "Excavate", "11", target_start_date="2026-02-02 08:00", target_end_date="2026-03-02 17:00",
           early_start_date="2026-02-02 08:00", early_end_date="2026-03-09 17:00",
           act_start_date="2026-02-02 08:00", act_end_date="2026-03-09 17:00",
           phys_complete_pct="100", target_drtn_hr_cnt="160"),
        _t("A200", "Pour L3 slab", "13", target_start_date="2026-09-01 08:00", target_end_date="2026-09-15 17:00",
           early_start_date="2026-09-10 08:00", early_end_date="2026-09-30 17:00", phys_complete_pct="40",
           act_start_date="2026-09-10 08:00", target_drtn_hr_cnt="80"),
        _t("M1", "Topping out", "12", task_type="TT_FinMile", target_end_date="2026-12-01 17:00",
           early_end_date="2026-12-01 17:00", target_drtn_hr_cnt="0"),
    ])
    doc = parse_schedule("sched.xer", data)
    assert doc["format"] == "xer" and doc["confidence"] == "structured"
    assert doc["data_date"] == "2026-10-05"  # last_recalc_date, not plan_start_date
    assert doc["project_name"] == "Tower 12"
    a = _by_id(doc)
    assert a["A100"]["baseline_finish"] == "2026-03-02"
    assert a["A100"]["current_finish"] == "2026-03-09"
    assert a["A100"]["actual_finish"] == "2026-03-09"
    assert a["A100"]["percent_complete"] == 100.0
    assert a["A100"]["duration_days"] == 20.0
    assert a["M1"]["kind"] == "milestone"
    assert "P6 target dates used as the original plan" in doc["warnings"]
    # hierarchy: root -> Superstructure -> Level 3 -> task
    assert a["WBS:10"]["kind"] == "summary" and a["WBS:10"]["level"] == 0
    assert a["WBS:13"]["level"] == 2 and a["A200"]["level"] == 3
    assert a["A200"]["wbs"] == "SUP.L3"
    assert "WBS:14" not in a  # a WBS node with nothing under it is not a row
    assert a["WBS:10"]["percent_complete"] is None  # P6 stores none on a WBS node


def test_xer_percent_only_when_physical():
    data = _xer([
        _t("D1", "Duration based", "11", complete_pct_type="CP_Drtn", phys_complete_pct="0",
           early_start_date="2026-02-02 08:00"),
        _t("P1", "Physical", "11", phys_complete_pct="25", early_start_date="2026-02-02 08:00"),
        _t("P2", "Physical empty", "11", phys_complete_pct="", early_start_date="2026-02-02 08:00"),
    ])
    doc = parse_schedule("a.xer", data)
    a = _by_id(doc)
    assert a["D1"]["percent_complete"] is None  # never the default 0 of a non-physical row
    assert a["P1"]["percent_complete"] == 25.0
    assert a["P2"]["percent_complete"] is None
    assert any("duration- or units-based" in w for w in doc["warnings"])


def test_xer_no_target_means_no_baseline_and_no_data_date():
    data = _xer([_t("A1", "Task", "11", early_start_date="2026-02-02 08:00", early_end_date="2026-02-09 08:00")],
                with_proj=False)
    doc = parse_schedule("a.xer", data)
    a = doc["activities"][-1]
    assert a["baseline_start"] is None and a["baseline_finish"] is None  # current never copied
    assert doc["data_date"] is None
    assert not any("target dates" in w for w in doc["warnings"])


def test_xer_sniffed_without_extension():
    data = _xer([_t("A1", "Task", "11", early_start_date="2026-02-02 08:00")])
    assert parse_schedule("export.dat", data)["format"] == "xer"


# ---------------------------------------------------------------- MS Project XML

_NS = "http://schemas.microsoft.com/project"


def _msp(tasks_xml, status="<StatusDate>2026-10-05T17:00:00</StatusDate>"):
    return (
        f'<?xml version="1.0" encoding="UTF-8"?><Project xmlns="{_NS}"><Name>t.xml</Name>'
        f"<Title>Tower 12</Title><MinutesPerDay>480</MinutesPerDay>{status}<Tasks>{tasks_xml}</Tasks></Project>"
    ).encode()


_MSP_TASKS = """
<Task><UID>0</UID><ID>0</ID><Name>Tower 12</Name><OutlineLevel>0</OutlineLevel><Summary>1</Summary>
 <Start>2026-02-02T08:00:00</Start><Finish>2026-12-01T17:00:00</Finish><PhysicalPercentComplete>35</PhysicalPercentComplete>
 <PercentComplete>90</PercentComplete>
 <Baseline><Number>0</Number><Start>2026-02-02T08:00:00</Start><Finish>2026-11-02T17:00:00</Finish></Baseline></Task>
<Task><UID>1</UID><ID>1</ID><Name>Foundations</Name><WBS>1</WBS><OutlineLevel>1</OutlineLevel><Summary>1</Summary>
 <Start>2026-02-02T08:00:00</Start><Finish>2026-04-02T17:00:00</Finish></Task>
<Task><UID>2</UID><ID>2</ID><Name>Pour footings</Name><WBS>1.1</WBS><OutlineLevel>2</OutlineLevel>
 <Start>2026-02-02T08:00:00</Start><Finish>2026-02-13T17:00:00</Finish><ActualStart>2026-02-03T08:00:00</ActualStart>
 <ActualFinish>2026-02-16T17:00:00</ActualFinish><PercentComplete>100</PercentComplete><Duration>PT80H0M0S</Duration>
 <Baseline><Number>0</Number><Start>2026-02-02T08:00:00</Start><Finish>2026-02-13T17:00:00</Finish></Baseline>
 <Baseline><Number>1</Number><Start>2026-03-01T08:00:00</Start><Finish>2026-03-13T17:00:00</Finish></Baseline></Task>
<Task><UID>3</UID><ID>3</ID><Name>Handover</Name><WBS>1.2</WBS><OutlineLevel>2</OutlineLevel><Milestone>1</Milestone>
 <Start>2026-12-01T08:00:00</Start><Finish>2026-12-01T08:00:00</Finish><Duration>PT0H0M0S</Duration></Task>
<Task><UID>4</UID><Null>1</Null><ID>4</ID></Task>
"""


def test_msp_xml_happy_path():
    doc = parse_schedule("t.xml", _msp(_MSP_TASKS))
    assert doc["format"] == "msp_xml" and doc["data_date"] == "2026-10-05"
    assert doc["project_name"] == "Tower 12"
    a = _by_id(doc)
    assert len(a) == 4  # the Null task is dropped
    assert a["0"]["kind"] == "summary" and a["0"]["level"] == 0 and a["0"]["percent_complete"] == 35.0
    assert a["2"]["level"] == 2 and a["2"]["wbs"] == "1.1"
    assert a["2"]["baseline_finish"] == "2026-02-13"  # Baseline Number 0 only
    assert a["2"]["actual_finish"] == "2026-02-16" and a["2"]["duration_days"] == 10.0
    assert a["3"]["kind"] == "milestone"
    assert a["3"]["baseline_start"] is None  # not baselined: not filled from current
    assert a["1"]["percent_complete"] is None  # not reported, not derived
    # NEGATIVE CONTROL: MS Project's PercentComplete is duration-based. Reading
    # it as completion would print "100% complete" for a task that merely ran
    # its planned duration; only PhysicalPercentComplete is a reported progress.
    assert a["2"]["percent_complete"] is None
    assert any("duration-based" in w for w in doc["warnings"])


def test_msp_xml_without_status_date_has_no_data_date():
    doc = parse_schedule("t.xml", _msp(_MSP_TASKS, status=""))
    assert doc["data_date"] is None
    assert any("status date" in w for w in doc["warnings"])


def test_xml_with_doctype_or_entities_is_refused():
    evil = (
        '<?xml version="1.0"?><!DOCTYPE p [<!ENTITY a "aaaa"><!ENTITY b "&a;&a;&a;&a;">]>'
        f'<Project xmlns="{_NS}"><Name>&b;</Name><Tasks/></Project>'
    ).encode()
    with pytest.raises(ScheduleImportError, match="DOCTYPE"):
        parse_schedule("t.xml", evil)
    with pytest.raises(ScheduleImportError, match="DOCTYPE"):
        parse_schedule("t.xml", evil.decode().encode("utf-16"))


def test_other_xml_is_refused():
    with pytest.raises(ScheduleImportError, match="MS Project"):
        parse_schedule("x.xml", b'<?xml version="1.0"?><root><Tasks><Task/></Tasks></root>')
    with pytest.raises(ScheduleImportError):
        parse_schedule("x.xml", b"<root><unclosed>")


# ---------------------------------------------------------------- Excel

def _xlsx(rows, *, sheets=None, formats=None, indents=None):
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl.styles import Alignment

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Schedule"
    for r in rows:
        ws.append(r)
    for coord, fmt in (formats or {}).items():
        ws[coord].number_format = fmt
    for coord, n in (indents or {}).items():
        ws[coord].alignment = Alignment(indent=n)
    for title, extra in (sheets or {}).items():
        w2 = wb.create_sheet(title)
        for r in extra:
            w2.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_happy_path_native_dates_baseline_and_actuals():
    data = _xlsx([
        ["Project: Tower 12"],
        ["Data Date", datetime(2026, 10, 5)],
        [],
        ["Task Name", "Start", "Finish", "Baseline Start", "Baseline Finish", "Actual Start", "Actual Finish", "% Complete"],
        ["Excavation", datetime(2026, 2, 2), datetime(2026, 3, 9), datetime(2026, 2, 2), datetime(2026, 3, 2),
         datetime(2026, 2, 3), datetime(2026, 3, 9), 100],
        ["Slab L3", datetime(2026, 9, 10), datetime(2026, 9, 30), None, None, datetime(2026, 9, 10), None, 40],
    ])
    doc = parse_schedule("plan.xlsx", data)
    assert doc["format"] == "xlsx" and doc["data_date"] == "2026-10-05" and doc["project_name"] == "Tower 12"
    a = _by_id(doc)
    ex = a["row1"]
    assert (ex["current_finish"], ex["baseline_finish"], ex["actual_finish"]) == ("2026-03-09", "2026-03-02", "2026-03-09")
    slab = a["row2"]
    assert slab["baseline_start"] is None and slab["baseline_finish"] is None  # never copied from current
    assert slab["percent_complete"] == 40.0 and slab["actual_finish"] is None
    assert any("had no id" in w for w in doc["warnings"])


def test_xlsx_header_synonyms_and_no_baseline_columns():
    data = _xlsx([
        ["Activity ID", "Activity", "Planned Start", "Planned Finish", "Progress"],
        ["A1", "Footings", "2026-02-02", "2026-02-13", "50%"],
        ["A2", "Walls", "2026-02-16", "2026-03-13", None],
    ])
    doc = parse_schedule("p.xlsx", data)
    a = _by_id(doc)
    assert a["A1"]["current_start"] == "2026-02-02" and a["A1"]["baseline_start"] is None
    assert a["A1"]["percent_complete"] == 50.0 and a["A2"]["percent_complete"] is None


def test_xlsx_percent_formatted_cells_and_plain_fraction_warning():
    data = _xlsx(
        [["Name", "Start", "Finish", "% Complete"],
         ["A", datetime(2026, 2, 2), datetime(2026, 2, 9), 0.5],
         ["B", datetime(2026, 2, 2), datetime(2026, 2, 9), 0.25]],
        formats={"D2": "0%", "D3": "0%"},
    )
    a = _by_id(parse_schedule("p.xlsx", data))
    assert a["row1"]["percent_complete"] == 50.0 and a["row2"]["percent_complete"] == 25.0
    plain = _xlsx([["Name", "Start", "Finish", "% Complete"], ["A", datetime(2026, 2, 2), datetime(2026, 2, 9), 0.5]])
    doc = parse_schedule("p.xlsx", plain)
    assert doc["activities"][0]["percent_complete"] == 0.5  # as written, never rescaled
    assert any("all 1 or less" in w for w in doc["warnings"])


def test_xlsx_percent_non_numbers_and_clamp():
    data = _xlsx([["Name", "Start", "Finish", "% Complete"],
                  ["A", "2026-02-02", "2026-02-09", "n/a"], ["B", "2026-02-02", "2026-02-09", 140],
                  ["C", "2026-02-02", "2026-02-09", True]])
    a = _by_id(parse_schedule("p.xlsx", data))
    assert a["row1"]["percent_complete"] is None
    assert a["row2"]["percent_complete"] == 100.0
    assert a["row3"]["percent_complete"] is None


def test_xlsx_levels_from_indent_and_summary_inference():
    data = _xlsx(
        [["Task Name", "Start", "Finish"],
         ["Phase 1", None, None],
         ["Dig", datetime(2026, 2, 2), datetime(2026, 2, 9)],
         ["Pour", datetime(2026, 2, 10), datetime(2026, 2, 20)],
         ["Phase 2", None, None],
         ["Frame", datetime(2026, 3, 2), datetime(2026, 3, 9)]],
        indents={"A3": 1, "A4": 1, "A6": 1},
    )
    doc = parse_schedule("p.xlsx", data)
    a = _by_id(doc)
    assert [(x["name"], x["level"], x["kind"]) for x in doc["activities"]] == [
        ("Phase 1", 0, "summary"), ("Dig", 1, "task"), ("Pour", 1, "task"),
        ("Phase 2", 0, "summary"), ("Frame", 1, "task"),
    ]
    assert a["row1"]["percent_complete"] is None


def test_outline_level_column_and_leading_spaces():
    csv_text = "ID,Task Name,Outline Level,Start,Finish\n1,Building,1,2026-02-02,2026-04-01\n2,Walls,2,2026-02-02,2026-03-01\n"
    doc = parse_schedule("a.csv", csv_text.encode())
    assert [(a["level"], a["kind"]) for a in doc["activities"]] == [(0, "summary"), (1, "task")]
    spaced = "Task,Start,Finish\nPhase,,\n    Dig,2026-02-02,2026-02-09\n"
    doc = parse_schedule("a.csv", spaced.encode())
    assert [(a["name"], a["level"]) for a in doc["activities"]] == [("Phase", 0), ("Dig", 1)]


def test_xlsx_picks_sheet_with_most_recognisable_rows():
    data = _xlsx(
        [["Name", "Start", "Finish"], ["only", "2026-02-02", "2026-02-09"]],
        sheets={"Full": [["Name", "Start", "Finish"], ["a", "2026-02-02", "2026-02-09"],
                         ["b", "2026-02-02", "2026-02-09"], ["c", "2026-02-02", "2026-02-09"]]},
    )
    doc = parse_schedule("p.xlsx", data)
    assert len(doc["activities"]) == 3
    assert any('"Full"' in w for w in doc["warnings"])


def test_xlsx_formulas_are_not_evaluated():
    openpyxl = pytest.importorskip("openpyxl")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "Start", "Finish"])
    ws.append(["a", "2026-02-02", "=DATE(2026,3,1)"])  # saved by openpyxl: no cached value
    buf = io.BytesIO()
    wb.save(buf)
    doc = parse_schedule("p.xlsx", buf.getvalue())
    assert doc["activities"][0]["current_finish"] is None


def test_xlsx_zip_bomb_and_garbage_refused(monkeypatch):
    data = _xlsx([["Name", "Start", "Finish"], ["a", "2026-02-02", "2026-02-09"]])
    monkeypatch.setattr(S, "MAX_XLSX_UNCOMPRESSED", 100)
    with pytest.raises(ScheduleImportError, match="expands"):
        parse_schedule("p.xlsx", data)
    monkeypatch.undo()
    with pytest.raises(ScheduleImportError):
        parse_schedule("p.xlsx", b"PK\x03\x04 not really a zip")


def test_xlsx_without_schedule_columns_refused():
    data = _xlsx([["Item", "Qty"], ["Bolts", 4]])
    with pytest.raises(ScheduleImportError, match="No schedule table"):
        parse_schedule("p.xlsx", data)


# ---------------------------------------------------------------- CSV and dates

_CSV_HEAD = "Name,Start,Finish\n"


def test_csv_us_and_dmy_order_detection():
    us = parse_schedule("a.csv", (_CSV_HEAD + "a,10/5/26,10/20/26\nb,11/1/26,11/9/26\n").encode())
    assert us["activities"][0]["current_start"] == "2026-10-05"  # 10/20 proves month/day
    assert not any("could be day/month" in w for w in us["warnings"])
    dm = parse_schedule("a.csv", (_CSV_HEAD + "a,5/10/26,20/10/26\n").encode())
    assert dm["activities"][0]["current_start"] == "2026-10-05"  # 20/10 proves day/month
    assert dm["activities"][0]["current_finish"] == "2026-10-20"


def test_csv_ambiguous_dates_assume_us_with_warning():
    doc = parse_schedule("a.csv", (_CSV_HEAD + "a,3/4/26,5/6/26\n").encode())
    assert doc["activities"][0]["current_start"] == "2026-03-04"
    assert any("could be day/month or month/day" in w for w in doc["warnings"])


def test_csv_text_date_forms_and_unreadable_dates():
    body = _CSV_HEAD + "a,Mon 10/5/26,\"Mon Oct 12, 26\"\nb,5-Oct-26,soon\nc,2026-10-05 08:00,5 October 2026\n"
    doc = parse_schedule("a.csv", body.encode())
    a = doc["activities"]
    assert (a[0]["current_start"], a[0]["current_finish"]) == ("2026-10-05", "2026-10-12")
    assert a[1]["current_start"] == "2026-10-05" and a[1]["current_finish"] is None
    assert a[2]["current_finish"] == "2026-10-05"
    assert any("could not be read as dates" in w for w in doc["warnings"])


def test_csv_encodings_delimiters_and_cleaning():
    latin = "Name;Start;Finish\nCaf\xe9\tbar\x07  slab;2026-02-02;2026-02-09\n".encode("latin-1")
    doc = parse_schedule("a.csv", latin)
    assert doc["activities"][0]["name"] == "Caf\xe9 bar slab"
    bom = "﻿Name\tStart\tFinish\nx\t2026-02-02\t2026-02-09\n".encode("utf-8")
    assert parse_schedule("a.tsv", bom)["activities"][0]["name"] == "x"


def test_csv_milestones_duplicates_and_empty_rows():
    body = (
        "ID,Name,Start,Finish,Duration,Milestone,% Complete,Actual Finish\n"
        "1,Pour,2026-02-02,2026-02-09,5 days,,50,2026-02-09\n"
        ",,,,,,,\n"
        "2,Permit,2026-02-02,2026-02-02,0 days,,,\n"
        "3,Flagged,2026-03-01,2026-03-05,3 days,Y,,\n"
        "3,Flagged,2026-03-01,2026-03-05,3 days,Y,,\n"
        "4,Zero but dates differ,2026-03-01,2026-03-05,0,,,\n"
    )
    doc = parse_schedule("a.csv", body.encode())
    a = _by_id(doc)
    assert len(doc["activities"]) == 4
    assert a["1"]["kind"] == "task" and a["1"]["duration_days"] == 5.0
    assert a["2"]["kind"] == "milestone" and a["3"]["kind"] == "milestone"
    assert a["4"]["kind"] == "task"
    assert a["1"]["percent_complete"] == 50.0  # actual finish + 50% is left alone
    assert any("duplicate" in w for w in doc["warnings"])
    assert any("actual finish but a reported percent" in w for w in doc["warnings"])


def test_csv_unrecognised_columns_are_warned_not_guessed():
    body = "Name,Start,Finish,Late Start,Colour\nx,2026-02-02,2026-02-09,2026-03-01,red\n"
    doc = parse_schedule("a.csv", body.encode())
    assert doc["activities"][0]["current_start"] == "2026-02-02"
    assert any("Late Start" in w and "Colour" in w for w in doc["warnings"])


def test_non_schedule_inputs_raise():
    for name, data in [
        ("a.csv", b"just,some,words\n1,2,3\n"),
        ("notes.txt", b"hello world"),
        ("pic.png", b"\x89PNG\r\n\x1a\n" + b"\x00" * 50),
        ("plan.mpp", b"\xd0\xcf\x11\xe0" + b"\x00" * 100),
        ("empty.csv", b""),
        ("a.xer", b"ERMHDR\t20\n%T\tPROJECT\n%F\tproj_id\n%R\t1\n%E\n"),
    ]:
        with pytest.raises(ScheduleImportError):
            parse_schedule(name, data)


def test_csv_start_only_is_allowed_and_dates_required():
    doc = parse_schedule("a.csv", b"Name,Start\nonly,2026-02-02\n")
    assert doc["activities"][0]["current_finish"] is None
    with pytest.raises(ScheduleImportError, match="readable date"):
        parse_schedule("a.csv", b"Name,Start,Finish\nonly,soon,later\n")


# ---------------------------------------------------------------- limits

def test_size_and_count_limits(monkeypatch):
    monkeypatch.setattr(S, "MAX_BYTES", 50)
    with pytest.raises(ScheduleImportError, match="larger than"):
        parse_schedule("a.csv", b"Name,Start,Finish\n" + b"x,2026-02-02,2026-02-09\n" * 5)
    monkeypatch.undo()
    monkeypatch.setattr(S, "MAX_ACTIVITIES", 3)
    rows = "".join(f"t{i},2026-02-02,2026-02-09\n" for i in range(4))
    with pytest.raises(ScheduleImportError, match="more than 3 activities"):
        parse_schedule("a.csv", (_CSV_HEAD + rows).encode())
    ok = "".join(f"t{i},2026-02-02,2026-02-09\n" for i in range(3))
    assert len(parse_schedule("a.csv", (_CSV_HEAD + ok).encode())["activities"]) == 3


def test_sha256_and_source_name():
    import hashlib

    data = (_CSV_HEAD + "a,2026-02-02,2026-02-09\n").encode()
    doc = parse_schedule("my  plan.csv", data)
    assert doc["sha256"] == hashlib.sha256(data).hexdigest()
    assert doc["source_name"] == "my plan.csv"


def test_parsing_never_touches_the_network(monkeypatch):
    def boom(*a, **k):
        raise AssertionError("network used")

    monkeypatch.setattr(socket, "socket", boom)
    monkeypatch.setattr(socket, "create_connection", boom)
    parse_schedule("a.csv", (_CSV_HEAD + "a,2026-02-02,2026-02-09\n").encode())
    parse_schedule("t.xml", _msp(_MSP_TASKS))
    parse_schedule("a.xer", _xer([_t("A1", "T", "11", early_start_date="2026-02-02 08:00")]))


# ---------------------------------------------------------------- PDF

def _pdf(rows, header=("ID", "Task Name", "Duration", "Start", "Finish"), pre=""):
    # importorskip lets an OSError (GTK missing) through as an error.
    if _weasyprint.SKIP:
        pytest.skip(_weasyprint.REASON)
    import weasyprint
    try:
        html = (
            "<html><body style='font-family:sans-serif;font-size:10pt'>" + pre
            + "<table style='border-collapse:collapse'>"
            + "<tr>" + "".join(f"<th style='border:1px solid #000;padding:3px'>{h}</th>" for h in header) + "</tr>"
            + "".join("<tr>" + "".join(f"<td style='border:1px solid #000;padding:3px'>{c}</td>" for c in r) + "</tr>" for r in rows)
            + "</table></body></html>"
        )
        return weasyprint.HTML(string=html).write_pdf()
    except Exception as e:  # pragma: no cover - native libs missing
        pytest.skip(f"weasyprint cannot render here: {e}")


def test_pdf_table_extracted_with_weekday_dates():
    pytest.importorskip("pdfplumber")
    rows = [(str(i), f"Task number {i}", "5 days", f"Mon 10/{i + 4}/26", f"Fri Oct {i + 8}, 26") for i in range(1, 8)]
    data = _pdf(rows, pre="<p>Project: Tower 12</p><p>Status Date: 10/5/26</p>")
    doc = parse_schedule("sched.pdf", data)
    assert doc["format"] == "pdf" and doc["confidence"] == "extracted"
    assert len(doc["activities"]) == 7
    a = _by_id(doc)
    assert a["1"]["current_start"] == "2026-10-05" and a["1"]["current_finish"] == "2026-10-09"
    assert a["1"]["baseline_start"] is None and a["1"]["percent_complete"] is None
    assert any("must be confirmed by a person" in w for w in doc["warnings"])
    assert doc["data_date"] == "2026-10-05"


def test_pdf_with_too_few_rows_or_no_table_refused():
    pytest.importorskip("pdfplumber")
    few = _pdf([("1", "Only", "5 days", "Mon 10/5/26", "Fri 10/9/26")] * 3)
    with pytest.raises(ScheduleImportError, match="does not contain a readable schedule table"):
        parse_schedule("s.pdf", few)
    nostart = _pdf([(str(i), f"T{i}", "5 days") for i in range(8)], header=("ID", "Task Name", "Duration"))
    with pytest.raises(ScheduleImportError, match="export to Excel/MS Project XML/P6 XER"):
        parse_schedule("s.pdf", nostart)
    with pytest.raises(ScheduleImportError, match="readable schedule table"):
        parse_schedule("s.pdf", b"%PDF-1.4 garbage")


def test_pdf_page_limit(monkeypatch):
    pytest.importorskip("pdfplumber")
    rows = [(str(i), f"Task {i}", "5 days", "10/5/26", "10/9/26") for i in range(6)]
    data = _pdf(rows)
    monkeypatch.setattr(S, "MAX_PDF_PAGES", 0)
    with pytest.raises(ScheduleImportError, match="more than 0 pages"):
        parse_schedule("s.pdf", data)
